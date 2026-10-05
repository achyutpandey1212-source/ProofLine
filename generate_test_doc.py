import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_proofline_test_doc():
    doc = Document()

    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles Setup
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x25, 0x2A)

    # Colors
    PRIMARY = RGBColor(0x0F, 0x3D, 0x5E)      # Deep Navy
    SECONDARY = RGBColor(0x02, 0x84, 0xC7)    # Cyan / Azure
    DARK_GRAY = RGBColor(0x37, 0x41, 0x51)
    ALERT_RED = RGBColor(0xB9, 0x1C, 0x1C)
    SUCCESS_GREEN = RGBColor(0x04, 0x78, 0x57)

    # ==========================================
    # TITLE & HEADER
    # ==========================================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("PROOFLINE: COMPREHENSIVE END-TO-END VERIFICATION & TEST SPECIFICATION")
    run_title.bold = True
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run("Full Application Quality Assurance, Edge Case Testing, Adversarial Scenarios & Crash Resilience Manual")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = SECONDARY
    run_sub.italic = True

    # Callout Box: Executive Scope
    tbl_callout = doc.add_table(rows=1, cols=1)
    tbl_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl_callout.rows[0].cells[0]
    set_cell_background(c, "F0F9FF")
    set_cell_margins(c, top=140, bottom=140, left=200, right=200)
    cp = c.paragraphs[0]
    cp.paragraph_format.space_after = Pt(4)
    r_hdr = cp.add_run("SYSTEM AUDIT OBJECTIVE & SCOPE\n")
    r_hdr.bold = True
    r_hdr.font.color.rgb = PRIMARY
    r_desc = cp.add_run(
        "This master test plan is engineered to validate ProofLine end-to-end. ProofLine is an automated chain-of-custody "
        "and physical recycling verification platform combining AI multi-modal document extraction (Gemini/ImageKit), "
        "deterministic reconciliation rules, interactive proof graphs, tamper-proof proof packets with Merkle hashing, "
        "adversarial simulation, and human-in-the-loop review. This manual provides exhaustive step-by-step test execution, "
        "pre-requisite materials, edge-case attacks ('against all odds'), and crash resilience verifications for zero failures."
    )
    r_desc.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================================
    # SECTION 1: MASTER INDEX & MATRIX
    # ==========================================
    h1 = doc.add_heading("1. MASTER TEST MATRIX & INDEX OF FEATURES", level=1)
    h1.style.font.color.rgb = PRIMARY

    idx_intro = doc.add_paragraph(
        "The following index outlines the 10 critical functional domains of ProofLine. Every module is verified across "
        "three testing dimensions: (1) Standard Happy-Path Flow, (2) Adversarial / 'Against Odds' Boundary Tests, and "
        "(3) Fault Tolerance & Crash Resilience."
    )
    idx_intro.paragraph_format.space_after = Pt(8)

    table_index = doc.add_table(rows=1, cols=4)
    table_index.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_index.autofit = False

    headers = ["Flow ID", "User Flow / Feature Name", "Primary Engine Components", "Test Artifacts Needed"]
    col_widths = [Inches(1.0), Inches(2.2), Inches(2.0), Inches(1.8)]

    hdr_row = table_index.rows[0]
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = col_widths[i]
        set_cell_background(cell, "0F3D5E")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)

    index_data = [
        ("UF-01", "Authentication & Access Control", "Firebase Auth, JWT, ProtectedRoute, KeyPool", "Valid/invalid Google & email accounts, expired JWTs"),
        ("UF-02", "Case Creation & Declaration", "CaseController, CaseValidator, Mongoose", "Partner names, Transaction IDs, Material types, Target weights"),
        ("UF-03", "Evidence Ingestion & Multi-File Upload", "Multer, ImageKit, SHA-256 Hashing, EvidenceModel", "Invoices, Scale Slips, Certificates, Corrupt/Oversized files"),
        ("UF-04", "AI Multi-Modal Data Extraction", "Gemini 1.5 Flash/Pro, Prompt Engineering, Fallback Pool", "Clear/blurry invoices, multi-angle scale displays, handwriting"),
        ("UF-05", "Deterministic Verification Engine", "14 Business Rules, Normalizer, Tolerance Logic", "Multi-image weight batches, non-standard units (tons/lbs/kg)"),
        ("UF-06", "Interactive Proof Graph & Node Lineage", "ProofGraphService, Canvas/Layout Engine, Drawer", "Complex multi-evidence cases, cyclic links, disconnected nodes"),
        ("UF-07", "Adversarial Attack Simulation Engine", "SimulationService, In-memory Clone, 6 Attack Scenarios", "Baseline verified case, weight fraud, duplicate tickets, cross-case files"),
        ("UF-08", "Human Review & Decision Resolution", "ReviewController, Case Lifecycle Status Machine", "Low-confidence extractions, discrepancy justification notes"),
        ("UF-09", "Audit-Ready Proof Packet & PDF Export", "ProofPacketService, Merkle Root Hasher, PDFKit", "Signed verification hash, QR codes, tampered hashes"),
        ("UF-10", "Public API v1 & Developer Playground", "ApiV1Controller, Idempotency, Rate Limiter, Sandboxes", "API Keys, Idempotency Keys, Postman/cURL, payload fuzzers")
    ]

    for row_idx, row_data in enumerate(index_data):
        row = table_index.add_row()
        bg_color = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            cell.width = col_widths[i]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = SECONDARY

    doc.add_page_break()

    # ==========================================
    # SECTION 2: DETAILED STEP-BY-STEP SPECIFICATIONS
    # ==========================================
    h2 = doc.add_heading("2. DETAILED TEST SPECIFICATIONS & CRASH RESILIENCE MANUAL", level=1)
    h2.style.font.color.rgb = PRIMARY

    flows = [
        {
            "id": "UF-01",
            "title": "Authentication, Session Lifecycle & Role Security",
            "materials": [
                "Authorized test account credentials (Email/Password & Google OAuth).",
                "Unauthorized / unregistered email credentials.",
                "Expired / Malformed Firebase JWT bearer tokens.",
                "Chrome / Firefox DevTools with Network throttling capabilities."
            ],
            "step_guide": [
                "1. Navigate to the landing page ('/') and click 'Sign In' or 'Launch App'.",
                "2. Submit valid test credentials on '/login'; confirm redirect to '/cases' with user avatar displayed in Header.",
                "3. Inspect browser localStorage and sessionStorage to confirm secure storage of Firebase user tokens.",
                "4. Refresh the page on protected routes ('/cases', '/cases/new', '/cases/:id/verify'); verify session is maintained without flashing unauthenticated screens.",
                "5. Click 'Sign Out'; ensure user session is purged, local state cleared, and immediately navigated to '/login'."
            ],
            "examination": [
                "Verify ProtectedRoute blocks all unauthenticated direct URL visits and redirects to '/login' while preserving redirect query parameters.",
                "Examine network traffic on '/api/cases' ensuring 'Authorization: Bearer <token>' header is automatically attached to every request.",
                "Confirm that upon token expiry, the application automatically refreshes tokens or gracefully prompts re-authentication rather than freezing."
            ],
            "odds_tests": [
                "TAMPERED JWT: Intercept request and change the JWT signature payload by 1 character. Verify API responds with 401 Unauthorized and frontend redirects cleanly to Login.",
                "RAPID LOGOUT/LOGIN: Repeatedly trigger login and logout within 2 seconds. Verify no memory leaks, unmounted component state updates, or hanging React promises.",
                "MULTIPLE TABS: Open two browser tabs logged in. Log out in Tab 1, then trigger a case creation in Tab 2. Confirm Tab 2 handles session invalidation without application crash."
            ],
            "crash_resilience": [
                "SIMULATE FIREBASE AUTH OUTAGE: Block `identitytoolkit.googleapis.com` via DevTools request blocking. Submit login. Verify user receives 'Authentication service temporarily unavailable' error banner, with zero React error boundaries triggered.",
                "LOCAL STORAGE QUOTA FULL: Artificially fill localStorage to quota limit. Ensure application logs warning and handles state in-memory without throwing unhandled fatal exceptions."
            ]
        },
        {
            "id": "UF-02",
            "title": "Case Declaration & Baseline Metadata Registration",
            "materials": [
                "Standard recycling shipment records: Partner Name ('Apex Plastics Ltd'), Transaction ID ('TX-99201'), Material ('PET Flakes'), Target Weight ('5500 kg').",
                "Edge-case records: Extreme quantities (0.001 kg, 1,000,000 kg), multi-lingual company names ('緑のリサイクル Corp'), special characters.",
                "MongoDB database connection inspector (or Compass) for verification of persisted state."
            ],
            "step_guide": [
                "1. From Dashboard ('/cases'), click '+ New Verification Case' navigating to '/cases/new'.",
                "2. Fill in Case Declaration form: Partner/Facility Name, Transaction Reference ID, Target Material Category, and Declared Total Weight with unit (kg/lbs/MT).",
                "3. Click 'Create Case'; verify loading button state activates, preventing duplicate double-clicks.",
                "4. Confirm redirection to Case Detail Page ('/cases/:caseId') with lifecycle badge indicating 'DRAFT' or 'PENDING_EVIDENCE'.",
                "5. Confirm MongoDB persistence: Case record created with unique 'caseId' (e.g. 'PL-CAS-xxxx'), initialized findings array, and authenticated user association."
            ],
            "examination": [
                "Ensure client-side input validation rejects empty partner names, negative weights, and zero-weight inputs with clear inline error labels.",
                "Examine backend Zod schema validation (`case.validator.ts`) ensuring invalid payloads return 400 Bad Request with field-specific diagnostics.",
                "Verify audit log or case history records 'CASE_CREATED' event with accurate timestamp."
            ],
            "odds_tests": [
                "MASSIVE INPUT LENGTH: Input 5,000 characters into Partner Name and Transaction ID. Verify server rejects or safely truncates without database buffer overflows.",
                "SQL / NOSQL INJECTION PAYLOAD: Input `{\"$ne\": null}` or `' OR 1=1 --` into Transaction ID. Verify payload is sanitized and treated strictly as literal string without MongoDB injection vulnerability.",
                "RAPID MULTI-CLICK: Click 'Create Case' 10 times in rapid succession. Verify idempotency or button disabling creates exactly ONE case."
            ],
            "crash_resilience": [
                "MONGODB DISCONNECTION: Disconnect MongoDB service briefly and submit case. Verify backend returns structured 500 error with message 'Database service unavailable' and frontend shows alert banner without white-screen crash.",
                "NETWORK DROPOUT: Cut connection during submit. Verify frontend retains form fields in state so the user does not lose their entered data."
            ]
        },
        {
            "id": "UF-03",
            "title": "Evidence Ingestion, File Hashing & Multi-Upload Pipeline",
            "materials": [
                "Valid Documents: Genuine PDF Tax Invoice, PDF Recycling Certificate, Scale Ticket Slips (JPG/PNG).",
                "Weighbridge Images: High-resolution clear scale displays showing tare, gross, and net weights.",
                "Hostile Files: 0-byte blank file, corrupted PDF header, 50MB oversized video file renamed to `.pdf`, `.exe` binary disguised as `.jpg`.",
                "Identical Files: Two copies of the same image with different filenames to test duplicate SHA-256 detection."
            ],
            "step_guide": [
                "1. Open Case Detail page ('/cases/:caseId'); locate Evidence Upload Dropzone.",
                "2. Select Evidence Type: 'INVOICE', 'SCALE_IMAGE', 'CERTIFICATE', or 'RECEIPT'.",
                "3. Drag and drop the corresponding document file.",
                "4. Confirm upload progress bar, cloud upload to ImageKit / S3, and server-side computation of SHA-256 checksum.",
                "5. Verify evidence card appears in Evidence List with preview thumbnail, computed file hash, document type badge, and 'PENDING_EXTRACTION' status."
            ],
            "examination": [
                "Inspect network response of `POST /cases/:id/evidence` ensuring payload contains `evidenceId`, `fileHash`, `url`, and `mimeType`.",
                "Verify file hash is an authentic 64-character hexadecimal SHA-256 hash matching the local `sha256sum <file>` output.",
                "Verify that uploading multiple scale images aggregates them properly under the case evidence array without overwriting."
            ],
            "odds_tests": [
                "DUPLICATE FILE COLLISION: Upload the exact same scale image twice within the same case. Confirm system either warns of duplicate file or flags Rule 11 (EVIDENCE_REUSE_DETECTION).",
                "FILE EXTENSION SPOOFING: Take an executable script (`test.sh` or `test.exe`) and rename it to `invoice.pdf`. Upload file. Verify backend MIME-type inspector rejects with 400 'Unsupported file format'.",
                "ZERO-BYTE FILE: Upload an empty 0-byte file. Verify upload fails cleanly with 'File is empty' error."
            ],
            "crash_resilience": [
                "IMAGEKIT SERVICE TIMEOUT: Mock ImageKit failure (HTTP 504 / Bad Gateway). Verify backend catches error gracefully, deletes temporary multer file from disk, and returns clear error without process crash.",
                "MAX FILE SIZE EXCEEDED: Upload a 60MB file exceeding Multer 25MB limit. Verify HTTP 413 Payload Too Large is returned with user-friendly error notice."
            ]
        },
        {
            "id": "UF-04",
            "title": "AI Multi-Modal Data Extraction (Gemini 1.5 & Fallback KeyPool)",
            "materials": [
                "Invoices with varied date formats (DD/MM/YYYY, YYYY-MM-DD, 'October 4, 2026').",
                "Invoices with non-standard units (e.g. 5.5 Metric Tons, 12,125 lbs, 5,500 kg).",
                "Low-quality / skewed photos of digital weighing scale LED displays.",
                "Handwritten weighbridge slips with low optical clarity.",
                "Mocked exhausted Gemini API key to trigger key pool rotation."
            ],
            "step_guide": [
                "1. Click 'Extract Data' or 'Process Evidence' on an uploaded evidence document.",
                "2. Observe extraction status badge changing from 'PENDING' to 'PROCESSING'.",
                "3. Backend invokes Gemini extraction service (`geminiExtraction.service.ts`) with structured prompt schema.",
                "4. Verify extracted JSON metadata populates the UI: Document Date, Entity/Partner Name, Extracted Weight, Weight Unit, Transaction Reference, and Extraction Confidence Score.",
                "5. For Scale Images, verify extraction captures measured scale reading and unit.",
                "6. Verify extraction confidence gauge (e.g. 0.95 green, 0.65 amber)."
            ],
            "examination": [
                "Verify extracted weights are normalized or structured with original raw text preserved for human audit.",
                "Ensure confidence score between 0.00 and 1.00 is strictly enforced; extractions below 0.75 are flagged under Rule 9 (EXTRACTION_CONFIDENCE).",
                "Verify keyPool service rotates API keys automatically if rate limits (HTTP 429) are encountered."
            ],
            "odds_tests": [
                "COMPLETE NONSENSE IMAGE: Upload an image of a landscape or cartoon cat marked as 'SCALE_IMAGE'. Trigger extraction. Verify model returns low confidence (<0.3) or extraction failure flag rather than hallucinating fictitious weights.",
                "EXTREME GLARE / BLUR: Upload a weighbridge photo where numbers are illegible. Verify system marks status as 'UNCERTAIN' and requests manual reviewer entry.",
                "MULTI-CURRENCY / MULTI-LINE INVOICE: Upload complex commercial invoice with line items. Verify parser correctly extracts the target commodity line and total net quantity."
            ],
            "crash_resilience": [
                "AI SERVICE OUTAGE: Disconnect AI provider. Ensure extraction service times out after 15s with fallback error state without hanging the HTTP event loop.",
                "ALL API KEYS EXHAUSTED: Simulate 429 quota exhaustion across all keyPool slots. Verify system returns 503 'AI extraction temporarily backlogged; please enter values manually or retry shortly'."
            ]
        },
        {
            "id": "UF-05",
            "title": "Deterministic Verification Engine & 14 Business Rules",
            "materials": [
                "Case A (Perfect Match): Declared 5,000 kg; Invoice 5,000 kg; 3 Scale Slips (1,650 kg + 1,700 kg + 1,650 kg = 5,000 kg); Matching Partner; Matching TxID.",
                "Case B (Minor Tolerance): Scale total 4,960 kg vs 5,000 kg Invoice (0.8% difference, within 2% tolerance).",
                "Case C (Adversarial Discrepancy): Scale total 3,200 kg vs 5,000 kg Invoice (36% shortage, critical mismatch).",
                "Case D (Chronological Flaw): Scale Weighing timestamp occurs AFTER the certificate was already issued.",
                "Case E (Entity Mismatch): Invoice from 'GreenPlanet Recyclers' vs Case Partner declared as 'Apex Waste Ltd'."
            ],
            "step_guide": [
                "1. With evidence uploaded and extracted, click 'Run Verification Engine' or navigate to '/cases/:caseId/verify'.",
                "2. Observe sequential rule execution visualization in the verification console.",
                "3. Confirm execution of critical verification rules: EVIDENCE_COMPLETENESS, EXTRACTION_CONFIDENCE, SCALE_AGGREGATION, DOCUMENT_QUANTITY_CONSISTENCY, WEIGHT_RECONCILIATION, ENTITY_CONSISTENCY, TRANSACTION_ID_CONSISTENCY, TEMPORAL_CONSISTENCY, SUSPICIOUS_UNIFORMITY_DETECTION.",
                "4. Check computed results: Calculated Total Measured Weight, Percentage Discrepancy, Overall Status ('PASS', 'REVIEW_REQUIRED', 'HIGH_RISK'), and List of Findings.",
                "5. Verify that findings include Finding ID, Severity ('LOW', 'MEDIUM', 'HIGH'), Involved Evidence IDs, Plain-Language Explanation, and Actionable Recommendation."
            ],
            "examination": [
                "Verify Rule 5 (Multi-Image Weight Aggregation) sums all SCALE_IMAGE evidence accurately before comparing.",
                "Verify Rule 6 (Weight Reconciliation) enforces the configurable 2.0% tolerance threshold.",
                "Ensure explainability rule: ProofLine never states 'Fraud detected'; it states 'High-risk evidence inconsistency detected' or 'Further investigation required'.",
                "Verify that missing non-critical documents return 'NOT_CHECKABLE' rather than false fails."
            ],
            "odds_tests": [
                "UNIT MISMATCH CONVERSION: Upload invoice in Metric Tons (5.0 MT) and scale slips in kg (5,000 kg). Verify normalizer converts units correctly so no false weight failure is triggered.",
                "IDENTICAL DECIMAL WEIGHTS: Provide three separate scale slips each claiming exactly '184.62 kg'. Verify Rule: SUSPICIOUS_UNIFORMITY_DETECTION flags suspicious measurement recycling.",
                "POST-DATED WEIGHING: Set scale ticket date to 2 days after the Certificate of Destruction. Verify Rule: TEMPORAL_CONSISTENCY triggers a chronological anomaly warning."
            ],
            "crash_resilience": [
                "NULL OR NAN VALUES: Submit evidence where AI returned null weight. Verify rule engine does not crash with `TypeError: Cannot read properties of null` or produce `NaN%` discrepancies.",
                "ZERO DIVISION GUARD: Submit declared weight = 0. Verify weight difference percentage calculation handles zero division gracefully without throwing `Infinity` or crashing the Node process."
            ]
        },
        {
            "id": "UF-06",
            "title": "Interactive Proof Graph & Node Lineage Visualization",
            "materials": [
                "Verified Case with minimum 4 connected evidence documents.",
                "Complex multi-scale transaction case.",
                "Modern browser with WebGL / HTML5 Canvas support enabled."
            ],
            "step_guide": [
                "1. From Case Detail or Verification Report page, click 'View Proof Graph' to launch the Graph Canvas modal.",
                "2. Inspect the rendered interactive graph: Case Root Node, Document Evidence Nodes, AI Extraction Nodes, Verification Rule Result Nodes, and Final Status Node.",
                "3. Click on individual nodes: verify GraphDetailDrawer slides out displaying deep metadata, confidence ratings, extracted attributes, and source file preview.",
                "4. Pan, zoom, and fit-to-screen using the canvas controls; confirm smooth 60fps rendering.",
                "5. If viewing a simulated case, confirm simulated nodes and mutated links are highlighted in neon-amber with 'SIMULATED' badges."
            ],
            "examination": [
                "Verify layoutEngine calculates non-overlapping coordinates for all nodes dynamically.",
                "Confirm that edge connectors show semantic relationship labels (e.g. 'VALIDATES', 'CONTAINS', 'EXTRACTED_FROM', 'DISCREPANCY_AGAINST').",
                "Verify responsive resizing when browser viewport is resized or inspected on tablet view."
            ],
            "odds_tests": [
                "DISCONNECTED / ORPHAN NODES: Test a case where one evidence item was not processed by AI. Verify graph engine renders orphan nodes gracefully without breaking topology.",
                "MASSIVE NODE GRAPH: Test a case with 25+ evidence items. Verify canvas layout engine does not cause browser UI freeze or memory spike.",
                "KEYBOARD & TOUCH ACCESSIBILITY: Test pinch-to-zoom on touch screens and Esc key to dismiss detail drawer."
            ],
            "crash_resilience": [
                "CORRUPT GRAPH DTO: Intercept API response of `/cases/:id/proof-graph` and pass invalid edge IDs referencing non-existent nodes. Verify frontend catches error and renders fallback 'Graph unavailable' card rather than crashing React."
            ]
        },
        {
            "id": "UF-07",
            "title": "Adversarial Simulation Engine (Zero Production Mutation)",
            "materials": [
                "A pristine verified Case with status 'PASS'.",
                "Simulation modal access on Case Detail page.",
                "MongoDB query tools to monitor production collections during simulation."
            ],
            "step_guide": [
                "1. Open a verified Case and click 'Simulate Attack Scenarios' or open SimulationModal.",
                "2. Select Scenario 1: 'WEIGHT_MISMATCH' (Mutates declared weight to trigger reconciliation failure).",
                "3. Click 'Run Adversarial Simulation'; verify SimulationBanner appears at top of screen in amber/black hazard styling: 'SIMULATION MODE ACTIVE: In-Memory Sandbox'.",
                "4. Inspect Verification Report under simulation: Status changes to 'REVIEW_REQUIRED' or 'HIGH_RISK'; finding FND-xxx highlights weight discrepancy.",
                "5. Open Proof Graph under simulation: Mutated node displays pulsing hazard ring and diff view showing `Original: 5000 kg -> Simulated: 7500 kg`.",
                "6. Test all 6 available simulation scenarios:\n   a. WEIGHT_MISMATCH\n   b. INVOICE_MISMATCH\n   c. TRANSACTION_MISMATCH\n   d. EVIDENCE_INCONSISTENCY\n   e. EVIDENCE_REUSE\n   f. CHRONOLOGY_ANOMALY.",
                "7. Exit simulation mode; confirm the page restores pristine production case data immediately."
            ],
            "examination": [
                "CRITICAL ZERO-MUTATION CHECK: Inspect MongoDB database before, during, and after simulation. Verify that CaseModel, EvidenceModel, and VerificationModel collections have ZERO modified records.",
                "Verify API endpoint `POST /cases/:caseId/simulate` executes purely in-memory using deep-cloned documents.",
                "Ensure simulation banner cannot be dismissed while viewing simulated data, preventing operators from mistaking simulation for real verification."
            ],
            "odds_tests": [
                "CONCURRENT SIMULATION: Two different browser sessions running two different simulation scenarios on the same case at the same time. Verify neither session interferes with the other.",
                "UNAUTHORIZED SIMULATION: Attempt to run simulation using an API key without sufficient permissions or on a case belonging to another user. Verify 403 Forbidden is returned."
            ],
            "crash_resilience": [
                "SIMULATION ON CASE WITH NO EVIDENCE: Attempt to trigger simulation on an empty Draft case. Verify backend returns 400 Bad Request 'No evidence files available to simulate' rather than crashing."
            ]
        },
        {
            "id": "UF-08",
            "title": "Human Review, Finding Override & Resolution Machine",
            "materials": [
                "A case in 'REVIEW_REQUIRED' status with 1 warning finding (e.g. 2.4% weight variance due to moisture loss).",
                "Authorized Reviewer / Compliance Officer credentials.",
                "Justification notes (e.g. 'Moisture loss verified via lab certificate')."
            ],
            "step_guide": [
                "1. Open Case Detail page for a case flagged for review; scroll to HumanReviewSection.",
                "2. Review each automated finding, its severity, and associated evidence files.",
                "3. Select Review Decision: 'APPROVE_WITH_OVERRIDE', 'REJECT_CLAIM', or 'REQUEST_MORE_EVIDENCE'.",
                "4. Enter mandatory Justification Note explaining the rationale.",
                "5. Submit review; verify Case Lifecycle status transitions to 'RESOLVED_APPROVED' or 'RESOLVED_REJECTED'.",
                "6. Confirm immutable audit trail: original automated rule results are permanently retained alongside reviewer comments, timestamp, and reviewer ID."
            ],
            "examination": [
                "Verify that submitting an override without entering a justification is blocked by frontend and backend validation.",
                "Verify that the system NEVER overwrites or deletes the original automated engine findings.",
                "Ensure that non-admin/unauthenticated users cannot submit review decisions."
            ],
            "odds_tests": [
                "CONCURRENT REVIEW SUBMISSION: Two reviewers submitting contradictory decisions simultaneously. Verify first submission commits and second receives conflict notification.",
                "SPECIAL CHARACTERS IN JUSTIFICATION: Submit emoji, markdown, and multilingual characters in review comments. Ensure safe rendering and storage without corruption."
            ],
            "crash_resilience": [
                "NETWORK DROPOUT ON SUBMIT: Disconnect connection during review post. Verify error prompt allows retry without losing entered explanation."
            ]
        },
        {
            "id": "UF-09",
            "title": "Audit-Ready Proof Packet & Tamper-Proof PDF Export",
            "materials": [
                "A fully verified and resolved case.",
                "PDF viewer (Adobe Acrobat, Chrome PDF Viewer).",
                "SHA-256 verification tool (OpenSSL, Python hashlib)."
            ],
            "step_guide": [
                "1. On Verification Report page ('/cases/:caseId/verification'), click 'Generate Proof Packet'.",
                "2. Inspect the generated cryptographic metadata: Merkle Root Hash, List of Leaf Hashes (Evidence files + Findings + Decisions), and Timestamped Digital Anchor.",
                "3. Click 'Download Audit PDF'; verify browser initiates download of `ProofLine-Verification-Packet-<CaseID>.pdf`.",
                "4. Open PDF: Verify professional vector layout, executive summary, verification breakdown table, embedded high-resolution QR Code, and SHA-256 proof footer.",
                "5. Scan embedded QR code using mobile phone camera; confirm it navigates to the public verification proof URL."
            ],
            "examination": [
                "Verify Merkle root is calculated deterministically: sorting hashes before hashing pairwise.",
                "Verify PDF generation runs via backend PDFKit service (`proofPacket.service.ts`) without blocking the Node event loop.",
                "Confirm PDF includes all findings, reviewer identity, and timestamp."
            ],
            "odds_tests": [
                "HASH TAMPERING TEST: Change 1 single character in any evidence file hash in the database. Regenerate Merkle root. Confirm the new root completely differs from original, proving tamper detection.",
                "EXPORT DURING SIMULATION: Ensure that if a user attempts to download a Proof Packet while in Simulation mode, the PDF is clearly watermarked 'SIMULATED DATA - NOT A VALID PROOF' or export is restricted."
            ],
            "crash_resilience": [
                "EXTREMELY LONG CASE HISTORY: Generate PDF for a case with 50+ scale evidence records. Verify PDFKit cleanly paginates across multi-page tables without clipping content or throwing memory heap errors."
            ]
        },
        {
            "id": "UF-10",
            "title": "Public REST API v1, Rate Limiting & Developer Playground",
            "materials": [
                "API Keys generated from '/cases' or '/playground'.",
                "Postman, Insomnia, or cURL command line terminal.",
                "Sandbox API key from `GET /api/v1/sandbox-key`.",
                "Stress-testing tool (Autocannon or loop script) to test rate limits."
            ],
            "step_guide": [
                "1. Navigate to Developer Playground ('/playground').",
                "2. Click 'Get Sandbox Key'; verify instant issuance of temporary test API key (`pl_live_...`).",
                "3. Execute cURL command to create verification:\n   `POST /api/v1/verifications` with Header `x-api-key: <key>` and `Idempotency-Key: <uuid>`.",
                "4. Upload evidence via API: `POST /api/v1/verifications/:id/evidence` multipart form.",
                "5. Run verification: `POST /api/v1/verifications/:id/run`.",
                "6. Fetch status and proof packet: `GET /api/v1/verifications/:id/proof-packet`.",
                "7. Verify interactive Playground UI synchronizes dynamically with live JSON responses and cURL snippets."
            ],
            "examination": [
                "Verify strict error contract: all API v1 errors follow `{ error: { code, message, requestId, details } }` format.",
                "Verify `X-Request-Id` is returned in response headers for traceability.",
                "Verify Rate Limiter: exceeds 60 requests per minute triggers HTTP 429 'Too Many Requests'."
            ],
            "odds_tests": [
                "IDEMPOTENCY RETRY: Send identical `POST /api/v1/verifications` request with same `Idempotency-Key` twice. Verify the second request returns identical HTTP 201 response cached from the first without creating a duplicate record in database.",
                "REVOKED API KEY: Revoke an API key from Dashboard. Immediately execute an API call using that key. Verify HTTP 401 Unauthorized 'Invalid or revoked API key'.",
                "MALFORMED JSON PAYLOAD: Send truncated JSON or invalid data types (e.g. `declaredWeight: 'five thousand'`). Verify HTTP 400 with Zod validation details."
            ],
            "crash_resilience": [
                "CONCURRENT SPAM: Send 120 rapid requests in 5 seconds. Confirm rate limiter throttles excess requests with 429 without service degradation or memory spikes on server."
            ]
        }
    ]

    for flow in flows:
        doc.add_heading(f"{flow['id']}: {flow['title']}", level=2)

        # Materials Callout
        p_mat_hdr = doc.add_paragraph()
        r_mh = p_mat_hdr.add_run("1. Pre-Requisite Test Materials & Environment Artifacts:")
        r_mh.bold = True
        r_mh.font.color.rgb = PRIMARY
        for mat in flow['materials']:
            p_m = doc.add_paragraph(style='List Bullet')
            p_m.paragraph_format.space_after = Pt(2)
            p_m.add_run(mat)

        # Step by Step Guide
        p_sg_hdr = doc.add_paragraph()
        p_sg_hdr.paragraph_format.space_before = Pt(6)
        r_sh = p_sg_hdr.add_run("2. Step-by-Step Execution & Examination Guide:")
        r_sh.bold = True
        r_sh.font.color.rgb = PRIMARY
        for step in flow['step_guide']:
            p_s = doc.add_paragraph()
            p_s.paragraph_format.space_after = Pt(3)
            p_s.paragraph_format.left_indent = Inches(0.2)
            p_s.add_run(step)

        # How to Examine & Verify
        p_ex_hdr = doc.add_paragraph()
        p_ex_hdr.paragraph_format.space_before = Pt(6)
        r_eh = p_ex_hdr.add_run("3. Examination & Success Verification Criteria:")
        r_eh.bold = True
        r_eh.font.color.rgb = SUCCESS_GREEN
        for exam in flow['examination']:
            p_e = doc.add_paragraph(style='List Bullet')
            p_e.paragraph_format.space_after = Pt(2)
            p_e.add_run(exam)

        # Against all Odds
        p_od_hdr = doc.add_paragraph()
        p_od_hdr.paragraph_format.space_before = Pt(6)
        r_oh = p_od_hdr.add_run("4. Testing Against Odds & Adversarial Scenarios:")
        r_oh.bold = True
        r_oh.font.color.rgb = ALERT_RED
        for odd in flow['odds_tests']:
            p_o = doc.add_paragraph(style='List Bullet')
            p_o.paragraph_format.space_after = Pt(2)
            p_o.add_run(odd)

        # Crash Resilience
        p_cr_hdr = doc.add_paragraph()
        p_cr_hdr.paragraph_format.space_before = Pt(6)
        r_ch = p_cr_hdr.add_run("5. Fault Tolerance & Crash Resilience Verification (Zero Unhandled Exceptions):")
        r_ch.bold = True
        r_ch.font.color.rgb = DARK_GRAY
        for cr in flow['crash_resilience']:
            p_c = doc.add_paragraph(style='List Bullet')
            p_c.paragraph_format.space_after = Pt(2)
            p_c.add_run(cr)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================================
    # SECTION 3: AUTOMATED TEST SUITE MAPPING
    # ==========================================
    doc.add_page_break()
    h3 = doc.add_heading("3. EXISTING TEST RUNNER SUITE & REPLICATION COMMANDS", level=1)
    h3.style.font.color.rgb = PRIMARY

    doc.add_paragraph(
        "ProofLine includes automated test suites covering backend integration, rule evaluation, API contracts, "
        "and adversarial simulation. Use the following commands to execute test suites and verify system baselines:"
    )

    tbl_tests = doc.add_table(rows=1, cols=3)
    tbl_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_headers = ["Test Suite File", "Covered Verification Domain", "CLI Command"]
    test_widths = [Inches(2.2), Inches(2.8), Inches(2.0)]
    for i, t_title in enumerate(test_headers):
        c = tbl_tests.rows[0].cells[i]
        c.width = test_widths[i]
        set_cell_background(c, "0F3D5E")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p = c.paragraphs[0]
        run = p.add_run(t_title)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    test_rows = [
        ("phase2.test.ts", "Authentication, User Registration, Key Pool Rotation", "npm run test:phase2"),
        ("phase3_1.test.ts", "Case Creation, Evidence Upload & ImageKit Mocking", "npm run test:phase3_1"),
        ("phase3_2.test.ts", "Gemini Multi-Modal Extraction & Validation", "npm run test:phase3_2"),
        ("phase3_3.test.ts", "Deterministic Rule Engine (14 Rules & Tolerances)", "npm run test:phase3_3"),
        ("phase3_4.test.ts", "Proof Graph Generation, Node Linking & Layout", "npm run test:phase3_4"),
        ("adversarial_simulator.test.ts", "In-Memory Attack Simulations (6 Scenarios)", "npx jest test/adversarial_simulator"),
        ("proof_packet.test.ts", "Merkle Root Calculation & PDFKit Generation", "npx jest test/proof_packet"),
        ("api_v1.test.ts", "External REST API v1, Idempotency & Rate Limiting", "npx jest test/api_v1"),
        ("review_workflow.test.ts", "Human Review & Lifecycle State Machine", "npx jest test/review_workflow")
    ]

    for r_idx, r_data in enumerate(test_rows):
        r = tbl_tests.add_row()
        bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(r_data):
            cell = r.cells[i]
            cell.width = test_widths[i]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r_val = p.add_run(val)
            r_val.font.size = Pt(9.5)
            if i == 0:
                r_val.bold = True
            elif i == 2:
                r_val.font.name = 'Consolas'
                r_val.font.size = Pt(8.5)

    # ==========================================
    # SECTION 4: QA SIGN-OFF CHECKLIST
    # ==========================================
    doc.add_page_break()
    h4 = doc.add_heading("4. END-TO-END QA SIGN-OFF CHECKLIST", level=1)
    h4.style.font.color.rgb = PRIMARY

    checklist_items = [
        ("Authentication & Sessions", "User can log in, sessions persist on refresh, logout clears all state, unauthorized routes redirect cleanly."),
        ("Case & Evidence Intake", "Declarations accept valid inputs, file upload calculates correct SHA-256 hashes, invalid extensions are rejected."),
        ("AI Extraction Resilience", "Extraction captures correct weights/dates, low-confidence scans trigger human review, API quota 429 fails gracefully."),
        ("Rule Reconciliation", "Tolerance thresholds (2%) work accurately, multi-scale images aggregate, discrepancies generate actionable findings."),
        ("Proof Graph Visualizer", "Nodes and links render without canvas overlap, detail drawer displays metadata, simulated nodes highlight accurately."),
        ("Adversarial Simulation", "All 6 attack scenarios trigger findings, MongoDB collections experience ZERO mutations, simulation banner alerts user."),
        ("Human Review Overrides", "Review decisions require justifications, original automated findings remain preserved, lifecycle state transitions properly."),
        ("Proof Packet & PDF", "Merkle root matches hash array, PDF downloads with clean vector formatting, QR code scans to valid URL."),
        ("Developer API v1", "API key authentication, Idempotency keys prevent double operations, Rate limit triggers 429 without server crashes.")
    ]

    tbl_chk = doc.add_table(rows=1, cols=3)
    tbl_chk.alignment = WD_TABLE_ALIGNMENT.CENTER
    chk_headers = ["Domain", "Verification Criteria", "QA Result (PASS/FAIL)"]
    chk_widths = [Inches(2.0), Inches(4.0), Inches(1.2)]
    for i, title in enumerate(chk_headers):
        c = tbl_chk.rows[0].cells[i]
        c.width = chk_widths[i]
        set_cell_background(c, "0F3D5E")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p = c.paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, (domain, crit) in enumerate(checklist_items):
        r = tbl_chk.add_row()
        bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([domain, crit, "[   ] PASS"]):
            cell = r.cells[i]
            cell.width = chk_widths[i]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(9.5)
            if i == 0:
                run.bold = True
            elif i == 2:
                run.bold = True
                run.font.color.rgb = DARK_GRAY

    # Save to disk
    output_filename = "ProofLine_Comprehensive_E2E_Test_Plan.docx"
    doc.save(output_filename)
    print(f"Document successfully saved as '{output_filename}' in {os.path.abspath(output_filename)}")

if __name__ == "__main__":
    create_proofline_test_doc()
